import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_body_p(doc, text, space_after=6, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def add_figure_entry(doc, fig_num_str, fig_title, img_path, desc_text):
    """
    Matches exact GRC Capstone Pages 67-87:
    Figure X. [Title]
    [IMAGE]
    Figure X [Description text...]
    """
    # Title above image
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(f"Figure {fig_num_str}. {fig_title}")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(12)
    r_title.bold = True
    
    # Image
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(6)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(5.8))
    
    # Description below image
    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc.paragraph_format.line_spacing = 1.5
    p_desc.paragraph_format.space_after = Pt(14)
    p_desc.paragraph_format.first_line_indent = Inches(0.5)
    
    # "Figure X " in bold or plain as in reference
    r_fig = p_desc.add_run(f"Figure {fig_num_str} ")
    r_fig.font.name = 'Times New Roman'
    r_fig.font.size = Pt(12)
    r_fig.bold = True
    
    r_text = p_desc.add_run(desc_text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(12)

def build_grc_chapter4(output_paths):
    doc = docx.Document()
    
    # Standard Letter 8.5 x 11, 1 inch margins
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    
    # Header
    p_chap = doc.add_paragraph()
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_chap.paragraph_format.space_after = Pt(2)
    p_chap.paragraph_format.line_spacing = 1.0
    r_chap = p_chap.add_run("CHAPTER IV")
    r_chap.bold = True
    r_chap.font.size = Pt(14)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(16)
    p_title.paragraph_format.line_spacing = 1.0
    r_title = p_title.add_run("RESULTS AND DISCUSSION")
    r_title.bold = True
    r_title.font.size = Pt(14)
    
    # Introductory Paragraph
    add_body_p(
        doc,
        "This chapter presents the results and detailed discussions of the web-based application development, "
        "focusing on how the project’s specific objectives were achieved and the effects of these outcomes. The "
        "chapter is arranged to match the specific objectives set at the start of this project, each covering an "
        "important part of the system’s functions."
    )
    
    # 4.1 Project Overview
    p_s41 = doc.add_paragraph()
    p_s41.paragraph_format.space_before = Pt(10)
    p_s41.paragraph_format.space_after = Pt(4)
    p_s41.paragraph_format.keep_with_next = True
    r_s41 = p_s41.add_run("4.1 Project Overview")
    r_s41.bold = True
    r_s41.font.size = Pt(12)
    
    add_body_p(
        doc,
        "The Web-Based Laundry Management System with Machine Learning for Demand Forecasting and Fleet Optimization "
        "is a centralized platform designed to improve the efficiency of operations for Bon Laundry Station, operating "
        "primarily for condominium residents of Urban Deca Homes Ortigas and walk-in patrons. It replaces manual, paper-based "
        "intake and uncertain walk-in waiting times with automated digital booking, transparent service catalogs, real-time "
        "machine occupancy telemetry, and live queue tracking."
    )
    
    add_body_p(
        doc,
        "The system provides dedicated role-based interfaces for Customers, Front-Desk Staff, and Station Managers. "
        "The customer-facing portal incorporates features such as a live 19-machine fleet lobby, structured service rate cards, "
        "a guided multi-machine load booking wizard, and live order tracking with queue countdowns. Through these modules, the "
        "system eliminates queue anxiety, prevents shop congestion, ensures transparent billing, and streamlines customer intake "
        "at the physical station counter."
    )
    
    # 4.2 Presentation of Results
    p_s42 = doc.add_paragraph()
    p_s42.paragraph_format.space_before = Pt(12)
    p_s42.paragraph_format.space_after = Pt(4)
    p_s42.paragraph_format.keep_with_next = True
    r_s42 = p_s42.add_run("4.2 Presentation of Results")
    r_s42.bold = True
    r_s42.font.size = Pt(12)
    
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(4)
    p_obj.paragraph_format.space_after = Pt(2)
    p_obj.paragraph_format.keep_with_next = True
    r_obj = p_obj.add_run("Specific Objective 1")
    r_obj.bold = True
    r_obj.font.size = Pt(12)
    
    p_bullet = doc.add_paragraph(style='List Bullet')
    p_bullet.paragraph_format.space_after = Pt(10)
    p_bullet.paragraph_format.line_spacing = 1.25
    r_b = p_bullet.add_run(
        "Develop an online customer portal with a live machine fleet lobby, official service rate cards, "
        "a multi-machine appointment booking wizard, and a real-time order tracking dashboard."
    )
    r_b.font.name = 'Times New Roman'
    r_b.font.size = Pt(12)
    
    # Figure 17 (or Figure 1)
    add_figure_entry(
        doc,
        "17",
        "Live Machine Fleet Lobby",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig1_fleet_lobby.jpg",
        "The Live Machine Fleet Lobby section provides the real-time operational status, capacity, and cycle tariff "
        "of all 19 commercial laundry machines at Bon Laundry Station. Customers can view the availability of 14 Regular "
        "Washers (8kg at ₱145 per cycle) and 5 Titan Dryers (13kg at ₱240 per cycle), with color-coded status badges indicating "
        "whether each unit is Available, In-Use, or under Maintenance. Cumulative run counters are displayed for each machine to "
        "support transparency in equipment maintenance and rotation. At the upper right, availability summary pills show total "
        "available units (19 Available, 0 In Cycle) alongside a manual sync button that allows customers to refresh the fleet status "
        "via client-side fetch without reloading the webpage. By providing live machine occupancy before arriving on-site, this "
        "module eliminates uncertainty, prevents shop floor congestion, and helps customers schedule their visits during off-peak hours."
    )
    
    # Figure 18 (or Figure 2)
    add_figure_entry(
        doc,
        "18",
        "Services & Official Rate Card",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig2_services_rates.jpg",
        "shows the Services and Official Rate Card module under the Customer Account Portal, which serves as the official "
        "commercial catalog of Bon Laundry Station. The module displays the available laundry packages, including Classic Care "
        "(₱130/8kg), Saver Package (₱160/8kg), Signature Package (₱200/8kg, highlighted as Best Value), and Premium Care (₱260/8kg "
        "for VIP garment separation), detailing all detergent, fabric conditioner, and folding inclusions. Below the package cards, "
        "turnaround options are presented, contrasting standard 24 Hours Express with a discounted 2 Days Economy option processed "
        "during graveyard shifts to optimize energy utilization. It also lists self-service wash/dry tariffs and the ₱40 flat-rate "
        "door-to-door condominium courier delivery fee for Urban Deca Homes residents. Clicking any \"Select & Book\" button or the "
        "\"Book Service →\" header link automatically pre-selects the package and redirects the customer to the booking wizard."
    )
    
    # Figure 19 (or Figure 3)
    add_figure_entry(
        doc,
        "19",
        "Book Laundry Appointment Wizard",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig3_booking_wizard.jpg",
        "presents the Book Laundry Appointment Wizard under the Customer Account Portal. This module provides a guided four-stage "
        "progress stepper (1. Service & Mode, 2. Schedule, 3. Contact & Pay, and 4. Review) that streamlines customer intake and "
        "eliminates manual errors. In Step 1, customers select their laundry package and configure multi-machine garment separation "
        "loads (1, 2, or 3+ loads) to isolate delicate garments, casual everyday clothes, or bulky bedsheets into separate 8kg or 13kg "
        "machines. Customers also designate their service mode: Counter Drop-Off, Self-Service Machine, or Deca Condo Delivery (+₱40), "
        "which dynamically expands building and unit fields in Step 2. Clicking \"Continue to Schedule →\" validates customer inputs and "
        "advances to succeeding steps for selecting guaranteed arrival windows, entering contact details for automated SMS updates, "
        "choosing payment methods, and generating a digital check-in ticket with an optical QR code for fast counter intake in under 15 seconds."
    )
    
    # Figure 20 (or Figure 4)
    add_figure_entry(
        doc,
        "20",
        "Active Orders & Live Queue Status",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig4_orders_tracking.jpg",
        "illustrates the Active Orders and Live Queue Status Tracking interface under the Customer Account Portal. This section "
        "enables customers to track their laundry orders in real time using their order code (e.g., ORD-66611) or by searching via the "
        "Search Order ID bar. The order card displays the scheduled service time, turnaround SLA, payment status (Pending or Paid), "
        "sequential queue token number (#82), and automated machine allocation status. A four-stage progress stepper visualizes the "
        "order's operational lifecycle: Stage 1 (Received / Intake & Weighed), Stage 2 (In Queue / Waiting Fleet Slot), Stage 3 "
        "(Processing / Wash & Dry Cycles), and Stage 4 (Ready / Folded & Packed). The interface also computes an estimated turnaround "
        "wait time (~65 mins) based on current queue depth and provides a \"Request Cancellation\" button that allows customers to cancel "
        "queued orders self-sufficiently, eliminating the need to physically call or visit the shop."
    )
    
    # 4.3 Testing Results
    p_s43 = doc.add_paragraph()
    p_s43.paragraph_format.space_before = Pt(14)
    p_s43.paragraph_format.space_after = Pt(4)
    p_s43.paragraph_format.keep_with_next = True
    r_s43 = p_s43.add_run("4.3 Testing Results")
    r_s43.bold = True
    r_s43.font.size = Pt(12)
    
    p_uit = doc.add_paragraph()
    p_uit.paragraph_format.space_before = Pt(2)
    p_uit.paragraph_format.space_after = Pt(6)
    p_uit.paragraph_format.keep_with_next = True
    r_uit = p_uit.add_run("Unit and Integration Testing")
    r_uit.bold = True
    r_uit.font.size = Pt(12)
    
    p_tbl_title = doc.add_paragraph()
    p_tbl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_title.paragraph_format.space_before = Pt(6)
    p_tbl_title.paragraph_format.space_after = Pt(4)
    p_tbl_title.paragraph_format.keep_with_next = True
    r_tbl_title = p_tbl_title.add_run("Table 24 Unit and Integration Testing for Web Application")
    r_tbl_title.italic = True
    r_tbl_title.bold = True
    r_tbl_title.font.size = Pt(11)
    
    headers = ["Test\nCase\nID", "Component/Module", "Test\nCondition", "Expected\nBehavior", "Actual Result", "Status"]
    table_data = [
        ("TC001", "Authentication", "User logs in with valid credentials, attempts login with incorrect password, and logs in using single sign-on authentication", "Valid credentials grant access to the appropriate role interface; invalid login is refused with no session created; single sign-on authenticates and redirects user", "Access granted for valid credentials; access refused for invalid credentials; single sign-on authenticated and redirected successfully", "Passed"),
        ("TC002", "Order Booking", "Customer places a laundry booking with selected service package, specifies condominium door-to-door delivery with missing and complete address details", "Incomplete address details are blocked with validation prompts; valid booking is saved with calculated service charges and delivery fees", "Missing address fields blocked; valid booking accepted and charges calculated correctly", "Passed"),
        ("TC003", "Order Verification", "Staff searches for an order using an order code, attempts lookup of an invalid code, and confirms garment intake", "Valid order details are retrieved and displayed; invalid code is rejected with an error notice; order confirmation updates the order status and moves it to the queue", "Order details retrieved correctly; invalid code rejected; order verified and moved to active queue", "Passed"),
        ("TC004", "Queue Management", "Verified orders enter the active queue, waiting times are calculated, and orders transition to processing upon machine allocation", "Queue numbers are issued sequentially; estimated waiting times compute based on machine load; queue status updates as orders progress", "Queue numbers issued in order; waiting times calculated accurately; status transitions updated without delay", "Passed"),
        ("TC005", "Machine Management", "Staff assigns laundry loads to available washers and dryers, attempts assignment of an overweight load, and starts machine cycle timers", "Machine status switches from available to in-use; overweight load is prevented with a capacity warning; cycle duration and usage count update", "Machine status updated to in-use; capacity limit enforced; cycle timers and usage count tracked accurately", "Passed"),
        ("TC006", "Payment Processing", "Customer selects cash or cashless payment methods, and staff updates cash payment status upon counter settlement", "Cash payment is marked pending; cashless payment is marked paid; manual status update marks order as paid", "Cash initialized as pending; cashless payment recorded as paid; manual update changed status to paid", "Passed"),
        ("TC007", "Order Tracking", "Customer tracks order progress across service stages, submits cancellation request, and verifies cancellation through a one-time verification code", "Current order stage and payment status are displayed; invalid verification code is rejected; valid verification code cancels order and releases queue position", "Order stages displayed accurately; invalid verification code rejected; valid verification code cancelled order and cleared queue", "Passed"),
        ("TC008", "Demand Forecasting", "Historical transaction records are processed by the regression model to forecast hourly customer demand", "Peak operating hours are projected; forecast accuracy is measured and recorded against acceptable performance limits", "Peak operating hours projected; forecast accuracy remained within acceptable limits", "Passed"),
        ("TC009", "Dashboard & Auditing", "Administrator accesses dashboard analytics, filters operational records, and performs administrative updates", "Daily statistics, machine utilization, and revenue totals compute correctly; record filters update dynamically; administrative actions are logged", "Dashboard values computed accurately; filter views updated dynamically; administrative actions recorded successfully", "Passed")
    ]
    
    table = doc.add_table(rows=len(table_data) + 2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    col_widths = [Inches(0.75), Inches(1.15), Inches(1.35), Inches(1.50), Inches(1.25), Inches(0.60)]
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=80, right=80)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))
    
    # Subheading row: "Web Application"
    sub_cells = table.rows[1].cells
    sub_cells[0].merge(sub_cells[5])
    sub_cells[0].text = "Web Application"
    p_sub = sub_cells[0].paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sub.paragraph_format.line_spacing = 1.0
    p_sub.paragraph_format.space_after = Pt(2)
    p_sub.paragraph_format.space_before = Pt(2)
    for run in p_sub.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.bold = True
    set_cell_background(sub_cells[0], "F2F2F2")
    set_cell_margins(sub_cells[0], top=50, bottom=50, left=80, right=80)
    
    # Data Rows
    for row_idx, data_tuple in enumerate(table_data, start=2):
        row_cells = table.rows[row_idx].cells
        for col_idx, text in enumerate(data_tuple):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if col_idx in [0, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                if col_idx == 0:
                    run.font.bold = True
            set_cell_margins(row_cells[col_idx], top=50, bottom=50, left=80, right=80)
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w
            
    # Summary paragraph
    add_body_p(
        doc,
        "Table 24 presents the unit and integration testing results for the Web Application. A total of nine (9) "
        "comprehensive test cases were executed across the primary functional components of the system, comprising "
        "Authentication, Order Booking, Order Verification, Queue Management, Machine Management, Payment Processing, "
        "Order Tracking, Demand Forecasting, and Dashboard & Auditing. All nine test cases achieved a status of Passed, "
        "demonstrating that each component operates reliably under both normal conditions and constraint-handling "
        "scenarios. The successful execution confirms that the web application satisfies its design requirements and "
        "provides dependable operational support for the laundry management workflow.",
        space_after=12
    )
    
    paths = [
        r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_DEVELOPMENT_AND_TESTING.docx',
        r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_PRESENTATION_OF_RESULTS.docx',
    ]
    try:
        paths.append(r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_RESULTS_AND_DISCUSSION.docx')
    except Exception:
        pass
    
    # Save safely
    for p in [r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_DEVELOPMENT_AND_TESTING.docx',
              r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_PRESENTATION_OF_RESULTS.docx']:
        doc.save(p)
        print(f"Generated clean GRC-style doc: {p}")
        
    try:
        doc.save(r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_RESULTS_AND_DISCUSSION.docx')
        print(f"Generated clean GRC-style doc: c:\\Users\\genes\\bon-laundry\\docs\\CHAPTER_4_RESULTS_AND_DISCUSSION.docx")
    except Exception as e:
        fallback = r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_RESULTS_AND_DISCUSSION_CLEAN.docx'
        doc.save(fallback)
        print(f"Note: CHAPTER_4_RESULTS_AND_DISCUSSION.docx was open in Word. Saved fallback to {fallback}")

if __name__ == '__main__':
    build_grc_chapter4([])

