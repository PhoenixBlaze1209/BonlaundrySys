import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_body_p(doc, text, space_after=6, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def add_figure_entry(doc, fig_num_str, fig_title, img_path, desc_text):
    """
    Matches GRC Capstone Pages 67-87:
    Figure X. [Title]
    [IMAGE]
    Figure X [Description text...]
    """
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(f"Figure {fig_num_str}. {fig_title}")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(12)
    r_title.bold = True
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(6)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(5.8))
    
    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc.paragraph_format.line_spacing = 1.5
    p_desc.paragraph_format.space_after = Pt(14)
    p_desc.paragraph_format.first_line_indent = Inches(0.5)
    
    r_fig = p_desc.add_run(f"Figure {fig_num_str} ")
    r_fig.font.name = 'Times New Roman'
    r_fig.font.size = Pt(12)
    r_fig.bold = True
    
    r_text = p_desc.add_run(desc_text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(12)

def build_full_grc_chapter4():
    doc = docx.Document()
    
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    
    # Header
    p_chap = doc.add_paragraph()
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_chap.paragraph_format.space_after = Pt(2)
    p_chap.paragraph_format.line_spacing = 1.0
    r_chap = p_chap.add_run("CHAPTER IV")
    r_chap.bold = True
    r_chap.font.size = Pt(14)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(16)
    p_title.paragraph_format.line_spacing = 1.0
    r_title = p_title.add_run("RESULTS AND DISCUSSION")
    r_title.bold = True
    r_title.font.size = Pt(14)
    
    # Introductory Paragraph
    add_body_p(
        doc,
        "This chapter presents the results and detailed discussions of the web-based application development, "
        "focusing on how the project’s specific objectives were achieved and the effects of these outcomes. The "
        "chapter is arranged to match the specific objectives set at the start of this project, each covering an "
        "important part of the system’s functions."
    )
    
    # 4.1 Project Overview
    p_s41 = doc.add_paragraph()
    p_s41.paragraph_format.space_before = Pt(10)
    p_s41.paragraph_format.space_after = Pt(4)
    p_s41.paragraph_format.keep_with_next = True
    r_s41 = p_s41.add_run("4.1 Project Overview")
    r_s41.bold = True
    r_s41.font.size = Pt(12)
    
    add_body_p(
        doc,
        "The Web-Based Laundry Management System with Machine Learning for Demand Forecasting and Fleet Optimization "
        "is a centralized platform designed to improve the efficiency of operations for Bon Laundry Station, operating "
        "primarily for condominium residents of Urban Deca Homes Ortigas and walk-in patrons. It replaces manual, paper-based "
        "intake and uncertain walk-in waiting times with automated digital booking, transparent service catalogs, real-time "
        "machine occupancy telemetry, and live queue tracking."
    )
    
    add_body_p(
        doc,
        "The system provides dedicated role-based interfaces for Customers, Front-Desk Staff, Station Managers, and Super Administrators. "
        "The customer-facing portal incorporates features such as a live 19-machine fleet lobby, structured service rate cards, "
        "a guided multi-machine load booking wizard, and live order tracking with queue countdowns. Correspondingly, the staff "
        "station console provides counter check-in order code verification, an appointments and condominium delivery lobby, live floor "
        "queue sequencing, machine cycle management, and machine learning store rush forecasts. For administrative governance, the manager "
        "console introduces operational audit logging, daily sales reconciliation, revenue tracking, and loss prevention ledgers, while "
        "the super administrator module enables instant role perspective switching and Google Single Sign-On account provisioning. "
        "Through these integrated modules, the system eliminates queue anxiety, prevents shop congestion, ensures transparent billing, "
        "and establishes complete managerial accountability."
    )
    
    # 4.2 Presentation of Results
    p_s42 = doc.add_paragraph()
    p_s42.paragraph_format.space_before = Pt(12)
    p_s42.paragraph_format.space_after = Pt(4)
    p_s42.paragraph_format.keep_with_next = True
    r_s42 = p_s42.add_run("4.2 Presentation of Results")
    r_s42.bold = True
    r_s42.font.size = Pt(12)
    
    # -------------------------------------------------------------
    # Specific Objective 1 (Customer Portal)
    # -------------------------------------------------------------
    p_obj1 = doc.add_paragraph()
    p_obj1.paragraph_format.space_before = Pt(4)
    p_obj1.paragraph_format.space_after = Pt(2)
    p_obj1.paragraph_format.keep_with_next = True
    r_obj1 = p_obj1.add_run("Specific Objective 1")
    r_obj1.bold = True
    r_obj1.font.size = Pt(12)
    
    p_b1 = doc.add_paragraph(style='List Bullet')
    p_b1.paragraph_format.space_after = Pt(10)
    p_b1.paragraph_format.line_spacing = 1.25
    r_b1_run = p_b1.add_run(
        "Develop an online customer portal with a live machine fleet lobby, official service rate cards, "
        "a multi-machine appointment booking wizard, and a real-time order tracking dashboard."
    )
    r_b1_run.font.name = 'Times New Roman'
    r_b1_run.font.size = Pt(12)
    
    # Figure 17
    add_figure_entry(
        doc,
        "17",
        "Customer Portal – Live Machine Fleet Lobby",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig1_fleet_lobby.jpg",
        "The Live Machine Fleet Lobby section provides the real-time operational status, capacity, and cycle tariff "
        "of all 19 commercial laundry machines at Bon Laundry Station. Customers can view the availability of 14 Regular "
        "Washers (8kg at ₱145 per cycle) and 5 Titan Dryers (13kg at ₱240 per cycle), with color-coded status badges indicating "
        "whether each unit is Available, In-Use, or under Maintenance. Cumulative run counters are displayed for each machine to "
        "support transparency in equipment maintenance and rotation. At the upper right, availability summary pills show total "
        "available units (19 Available, 0 In Cycle) alongside a manual sync button that allows customers to refresh the fleet status "
        "via client-side fetch without reloading the webpage. By providing live machine occupancy before arriving on-site, this "
        "module eliminates uncertainty, prevents shop floor congestion, and helps customers schedule their visits during off-peak hours."
    )
    
    # Figure 18
    add_figure_entry(
        doc,
        "18",
        "Customer Portal – Services & Official Rate Card",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig2_services_rates.jpg",
        "shows the Services and Official Rate Card module under the Customer Account Portal, which serves as the official "
        "commercial catalog of Bon Laundry Station. The module displays the available laundry packages, including Classic Care "
        "(₱130/8kg), Saver Package (₱160/8kg), Signature Package (₱200/8kg, highlighted as Best Value), and Premium Care (₱260/8kg "
        "for VIP garment separation), detailing all detergent, fabric conditioner, and folding inclusions. Below the package cards, "
        "turnaround options are presented, contrasting standard 24 Hours Express with a discounted 2 Days Economy option processed "
        "during graveyard shifts to optimize energy utilization. It also lists self-service wash/dry tariffs and the ₱40 flat-rate "
        "door-to-door condominium courier delivery fee for Urban Deca Homes residents. Clicking any \"Select & Book\" button or the "
        "\"Book Service →\" header link automatically pre-selects the package and redirects the customer to the booking wizard."
    )
    
    # Figure 19
    add_figure_entry(
        doc,
        "19",
        "Customer Portal – Book Laundry Appointment Wizard",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig3_booking_wizard.jpg",
        "presents the Book Laundry Appointment Wizard under the Customer Account Portal. This module provides a guided four-stage "
        "progress stepper (1. Service & Mode, 2. Schedule, 3. Contact & Pay, and 4. Review) that streamlines customer intake and "
        "eliminates manual errors. In Step 1, customers select their laundry package and configure multi-machine garment separation "
        "loads (1, 2, or 3+ loads) to isolate delicate garments, casual everyday clothes, or bulky bedsheets into separate 8kg or 13kg "
        "machines. Customers also designate their service mode: Counter Drop-Off, Self-Service Machine, or Deca Condo Delivery (+₱40), "
        "which dynamically expands building and unit fields in Step 2. Clicking \"Continue to Schedule →\" validates customer inputs and "
        "advances to succeeding steps for selecting guaranteed arrival windows, entering contact details for automated SMS updates, "
        "choosing payment methods, and generating a digital check-in ticket with an optical QR code for fast counter intake in under 15 seconds."
    )
    
    # Figure 20
    add_figure_entry(
        doc,
        "20",
        "Customer Portal – Active Orders & Live Queue Status",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig4_orders_tracking.jpg",
        "illustrates the Active Orders and Live Queue Status Tracking interface under the Customer Account Portal. This section "
        "enables customers to track their laundry orders in real time using their order code (e.g., ORD-66611) or by searching via the "
        "Search Order ID bar. The order card displays the scheduled service time, turnaround SLA, payment status (Pending or Paid), "
        "sequential queue token number (#82), and automated machine allocation status. A four-stage progress stepper visualizes the "
        "order's operational lifecycle: Stage 1 (Received / Intake & Weighed), Stage 2 (In Queue / Waiting Fleet Slot), Stage 3 "
        "(Processing / Wash & Dry Cycles), and Stage 4 (Ready / Folded & Packed). The interface also computes an estimated turnaround "
        "wait time (~65 mins) based on current queue depth and provides a \"Request Cancellation\" button that allows customers to cancel "
        "queued orders self-sufficiently, eliminating the need to physically call or visit the shop."
    )
    
    # -------------------------------------------------------------
    # Specific Objective 2 (Staff Station Console)
    # -------------------------------------------------------------
    p_obj2 = doc.add_paragraph()
    p_obj2.paragraph_format.space_before = Pt(14)
    p_obj2.paragraph_format.space_after = Pt(2)
    p_obj2.paragraph_format.keep_with_next = True
    r_obj2 = p_obj2.add_run("Specific Objective 2")
    r_obj2.bold = True
    r_obj2.font.size = Pt(12)
    
    p_b2 = doc.add_paragraph(style='List Bullet')
    p_b2.paragraph_format.space_after = Pt(10)
    p_b2.paragraph_format.line_spacing = 1.25
    r_b2_run = p_b2.add_run(
        "Implement a staff station console that streamlines counter check-in verification, coordinates condominium doorstep "
        "deliveries, manages floor queue dispatching, monitors machine cycle telemetry, and integrates machine learning store rush forecasts."
    )
    r_b2_run.font.name = 'Times New Roman'
    r_b2_run.font.size = Pt(12)
    
    # Figure 21 (Staff Dashboard)
    add_figure_entry(
        doc,
        "21",
        "Staff Console – Operations Dashboard",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig21_staff_dashboard.jpg",
        "shows the Operations Dashboard under the Staff Station Console. This interface serves as the primary operational command "
        "center for front-desk staff, providing real-time visibility into shop floor activity and predictive workload insights. At the top "
        "of the dashboard, four key performance indicator (KPI) cards display the Active Queue count (19 customers currently waiting), "
        "Estimated Turnaround Wait time (65 minutes), Ready Machines ratio (19/19 available for immediate wash or dry), and Today’s Workload "
        "status (Active Shift). On the left panel, the Live Machine Activity & Queue section lists waiting customer tokens with service "
        "tier badges (e.g., 24h Express) and payment indicators. On the right panel, the AI Store Rush Forecast module predicts daily "
        "customer traffic trends using machine learning regression, highlighting expected peak hours (e.g., 2:00 PM) to advise staff on "
        "break scheduling and resource preparation before rush hours occur."
    )
    
    # Figure 22 (Staff Verify Order)
    add_figure_entry(
        doc,
        "22",
        "Staff Console – Order Verification Terminal",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig22_staff_verify.jpg",
        "presents the Order Verification Terminal under the Staff Station Console. This module is utilized by front-desk personnel to "
        "process customer check-ins quickly when patrons arrive at the physical counter. The staff member scans or inputs the customer's "
        "unique Order Reference Code (e.g., ORD-66611) into the lookup field and clicks \"Look Up\". The terminal instantly retrieves the "
        "booking record from the database, displaying the customer’s name, mobile number, scheduled arrival slot, service package, service "
        "mode, turnaround SLA, and payment status (such as GCash or Pending). Staff can review the order details and click the green "
        "\"Verify & Queue Order\" button to confirm garment intake, which automatically transitions the order into the active shop queue "
        "and issues a sequential queue token. This barcode and code verification feature reduces intake registration time to under 15 "
        "seconds, preventing long lines at the front desk."
    )
    
    # Figure 23 (Staff Appointments Tracker)
    add_figure_entry(
        doc,
        "23",
        "Staff Console – Appointments Tracker and Delivery Lobby",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig23_staff_appointments.jpg",
        "illustrates the Appointments Tracker and Delivery Lobby module under the Staff Station Console. This section coordinates scheduled "
        "customer drop-offs and condominium doorstep delivery operations for residents of Urban Deca Homes. At the top of the lobby, "
        "summary cards categorize operational workloads by displaying Total Active orders (20), Condo Pickups (5), Counter Drop-Offs (15), "
        "Cash Pending collections (16), and Transit & Ready orders (1). Below the summary cards, staff can filter orders by service mode "
        "or cash collection status and search records by order ID, customer name, or phone number. Each appointment card displays the ticket "
        "number, customer contact details, scheduled drop-off window, payment badge, and designated building and unit address for "
        "condominium pickups. Action buttons such as \"Accept & Start →\", \"Collect ₱\", and \"Pick Up by Staff →\" allow operators to "
        "update logistics and payment states seamlessly as orders progress."
    )
    
    # Figure 24 (Staff Queue Order)
    add_figure_entry(
        doc,
        "24",
        "Staff Console – Live Queue Order and Floor Dispatch",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig24_staff_queue.jpg",
        "shows the Live Queue Order and Floor Dispatch interface under the Staff Station Console. This module manages the real-time "
        "sequencing of verified customer orders awaiting washer and dryer allocation. The upper telemetry header displays real-time line "
        "metrics, including Total in Line (19), Waiting Machine (19), In Wash/Dry Cycle (0), and Estimated Floor Wait time. A search bar "
        "and a \"+ Add Walk-In\" button enable operators to quickly register walk-in customers directly into the queue. In the dispatch table, "
        "each customer is listed in order of arrival with their queue position (#53, #54, etc.), reference code, payment status, and "
        "estimated wait duration. Staff can click the blue \"Start Machine\" floor action button to assign an order to an idle commercial "
        "washer, triggering the cycle timer and advancing the customer's order to the processing stage."
    )
    
    # Figure 25 (Staff Machine Fleet)
    add_figure_entry(
        doc,
        "25",
        "Staff Console – Machine Fleet Telemetry",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig25_staff_fleet.jpg",
        "depicts the Machine Fleet Telemetry module under the Staff Station Console. This screen provides staff with full control and "
        "oversight of all 19 commercial laundry machines (14 Regular 8kg Washers and 5 Titan 13kg Dryers) installed at the station. "
        "Real-time filter pills allow staff to toggle views between washers, dryers, available machines, and in-use units, while a search "
        "bar enables quick lookup of specific machine numbers. Each machine card indicates the unit’s capacity rating, fixed per-cycle "
        "tariff (₱145 for washers, ₱240 for Titan dryers), real-time operational status (Available in green, In-Use in blue, or Offline in red), "
        "and current queue occupancy. Staff can click \"+ Load Machine at Counter\" directly on any available machine card to bind an active "
        "counter order to that specific unit, automatically initiating the wash or dry cycle and synchronizing machine occupancy with "
        "customer-side telemetry."
    )
    
    # -------------------------------------------------------------
    # Specific Objective 3 (Manager Administration Console)
    # -------------------------------------------------------------
    p_obj3 = doc.add_paragraph()
    p_obj3.paragraph_format.space_before = Pt(14)
    p_obj3.paragraph_format.space_after = Pt(2)
    p_obj3.paragraph_format.keep_with_next = True
    r_obj3 = p_obj3.add_run("Specific Objective 3")
    r_obj3.bold = True
    r_obj3.font.size = Pt(12)
    
    p_b3 = doc.add_paragraph(style='List Bullet')
    p_b3.paragraph_format.space_after = Pt(10)
    p_b3.paragraph_format.line_spacing = 1.25
    r_b3_run = p_b3.add_run(
        "Establish a manager administration console that maintains operational audit logs, tracks multi-tier sales metrics, "
        "reconciles daily cash collections, and enforces accountability across staff operations."
    )
    r_b3_run.font.name = 'Times New Roman'
    r_b3_run.font.size = Pt(12)
    
    # Figure 26 (Manager Operational Audit Trail)
    add_figure_entry(
        doc,
        "26",
        "Manager Console – Operational Audit Trail",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig26_manager_audit.jpg",
        "shows the Operational Audit Trail under the Manager Console. This module provides station managers and administrators "
        "with comprehensive oversight of daily floor actions, accountability tracking, and security auditing. At the top of the interface, "
        "summary metrics indicate Events Today (3 actions), Orders Handled Today (0 verifications and queue entries), Machine Changes "
        "Today (0 state transitions), and Total Recorded Logs (75 active logs). The Staff Operational Activity Log records every key action "
        "executed in the system in chronological sequence, capturing the exact timestamp, actor name and role (e.g., Super Admin, Manager, "
        "Staff, Customer), event type (such as User Login, Role View Switched, or Payment Updated), detailed target information, and client "
        "IP address. Managers can filter the audit trail by event type or search specific order codes, staff members, and keywords. By "
        "maintaining an immutable log of counter payments, queue modifications, and role switches, this feature enforces operational "
        "transparency and prevents unauthorized system tampering."
    )
    
    # Figure 27 (Manager Sales History & Tally)
    add_figure_entry(
        doc,
        "27",
        "Manager Console – Sales History and Daily Tally",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig27_manager_sales.jpg",
        "illustrates the Sales History and Daily Tally module under the Manager Console, which serves as the financial ledger and "
        "loss-prevention tool for Bon Laundry Station. The upper telemetry cards summarize key revenue performance indicators, including "
        "Today's Audited Gross (₱250.00 from 1 checkout and 8kg load), Rolling 7-Day Volume (₱250.00), Current Month Revenue (₱29,315.00), "
        "and All-Time Revenue (₱72,218.00 across 370 total orders and 3,496.53kg processed). In the Daily Gross Tally & Reconciliation tab, "
        "daily financial transactions are aggregated to display order counts, total laundry weight, paid cash/e-wallet collected, pending "
        "receivables, gross totals, and reconciliation status (marked as Balanced). Managers can adjust the reporting timeframe using the "
        "View Days dropdown, inspect individual transactions via the Itemized Receipts Ledger, or click \"Export CSV\" to download financial "
        "records for external accounting and audits, effectively eliminating cash discrepancies between staff shifts."
    )
    
    # -------------------------------------------------------------
    # Specific Objective 4 (Super Admin System Governance & Account Management)
    # -------------------------------------------------------------
    p_obj4 = doc.add_paragraph()
    p_obj4.paragraph_format.space_before = Pt(14)
    p_obj4.paragraph_format.space_after = Pt(2)
    p_obj4.paragraph_format.keep_with_next = True
    r_obj4 = p_obj4.add_run("Specific Objective 4")
    r_obj4.bold = True
    r_obj4.font.size = Pt(12)
    
    p_b4 = doc.add_paragraph(style='List Bullet')
    p_b4.paragraph_format.space_after = Pt(10)
    p_b4.paragraph_format.line_spacing = 1.25
    r_b4_run = p_b4.add_run(
        "Implement a super administrator module that provides role-based perspective switching and automated single sign-on "
        "user account management for staff onboarding and role governance."
    )
    r_b4_run.font.name = 'Times New Roman'
    r_b4_run.font.size = Pt(12)
    
    # Figure 28 (Super Admin Switch View Mode)
    add_figure_entry(
        doc,
        "28",
        "Super Admin – Switch View Mode Perspective Selector",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig28_superadmin_switch_role.jpg",
        "shows the Switch View Mode Perspective Selector accessible to the Super Administrator from the Station Console header. "
        "This module enables administrators to switch dynamically between four system perspectives—Super Admin (full audit trail, governance, "
        "and settings), Manager (store operations, sales, and machine fleet), Staff (counter order verification and queue dispatch), and "
        "Customer Portal (self-service booking and order tracking)—without terminating the session or repeatedly logging in with different "
        "credentials. When an option is selected, the application updates the active session role in memory and adjusts the navigational "
        "sidebar and permissions instantaneously. By allowing administrators to audit the interface and test workflows from the exact "
        "perspective of each stakeholder group, this feature simplifies system verification, accelerates staff assistance, and enhances "
        "operational governance."
    )
    
    # Figure 29 (Super Admin User & Staff Account Management)
    add_figure_entry(
        doc,
        "29",
        "Super Admin – User and Staff Account Management",
        r"c:\Users\genes\bon-laundry\docs\screenshots\fig29_superadmin_account_mgmt.jpg",
        "illustrates the User and Staff Account Management interface located under System Settings, accessible exclusively to the Super "
        "Administrator. The upper Google SSO Pre-Registration section provides a secure form to create and onboard new employees without "
        "requiring manual password configuration. The administrator inputs the employee's Full Name, official Gmail address, and designates "
        "an Assigned System Role (such as Front-Desk Staff / Cashier or Manager) before clicking \"+ Pre-Register Staff Account\", allowing "
        "new hires to sign in immediately using Google Single Sign-On. Below the creation form, the User & Staff Directory provides a "
        "centralized table displaying all registered accounts, their current authorization role (Super Admin, Manager, Staff), and real-time "
        "Google OAuth connection state (Connected vs. Pending Sign-In). A \"Promote / Reassign\" dropdown enables administrators to modify "
        "permissions or reassign roles dynamically, ensuring strict access control and reliable employee lifecycle management."
    )
    
    # -------------------------------------------------------------
    # 4.3 Testing Results
    # -------------------------------------------------------------
    p_s43 = doc.add_paragraph()
    p_s43.paragraph_format.space_before = Pt(16)
    p_s43.paragraph_format.space_after = Pt(4)
    p_s43.paragraph_format.keep_with_next = True
    r_s43 = p_s43.add_run("4.3 Testing Results")
    r_s43.bold = True
    r_s43.font.size = Pt(12)
    
    p_uit = doc.add_paragraph()
    p_uit.paragraph_format.space_before = Pt(2)
    p_uit.paragraph_format.space_after = Pt(6)
    p_uit.paragraph_format.keep_with_next = True
    r_uit = p_uit.add_run("Unit and Integration Testing")
    r_uit.bold = True
    r_uit.font.size = Pt(12)
    
    p_tbl_title = doc.add_paragraph()
    p_tbl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_title.paragraph_format.space_before = Pt(6)
    p_tbl_title.paragraph_format.space_after = Pt(4)
    p_tbl_title.paragraph_format.keep_with_next = True
    r_tbl_title = p_tbl_title.add_run("Table 24 Unit and Integration Testing for Web Application")
    r_tbl_title.italic = True
    r_tbl_title.bold = True
    r_tbl_title.font.size = Pt(11)
    
    headers = ["Test\nCase\nID", "Component/Module", "Test\nCondition", "Expected\nBehavior", "Actual Result", "Status"]
    table_data = [
        ("TC001", "Authentication", "User logs in with valid credentials, attempts login with incorrect password, and logs in using single sign-on authentication", "Valid credentials grant access to the appropriate role interface; invalid login is refused with no session created; single sign-on authenticates and redirects user", "Access granted for valid credentials; access refused for invalid credentials; single sign-on authenticated and redirected successfully", "Passed"),
        ("TC002", "Order Booking", "Customer places a laundry booking with selected service package, specifies condominium door-to-door delivery with missing and complete address details", "Incomplete address details are blocked with validation prompts; valid booking is saved with calculated service charges and delivery fees", "Missing address fields blocked; valid booking accepted and charges calculated correctly", "Passed"),
        ("TC003", "Order Verification", "Staff searches for an order using an order code, attempts lookup of an invalid code, and confirms garment intake", "Valid order details are retrieved and displayed; invalid code is rejected with an error notice; order confirmation updates the order status and moves it to the queue", "Order details retrieved correctly; invalid code rejected; order verified and moved to active queue", "Passed"),
        ("TC004", "Queue Management", "Verified orders enter the active queue, waiting times are calculated, and orders transition to processing upon machine allocation", "Queue numbers are issued sequentially; estimated waiting times compute based on machine load; queue status updates as orders progress", "Queue numbers issued in order; waiting times calculated accurately; status transitions updated without delay", "Passed"),
        ("TC005", "Machine Management", "Staff assigns laundry loads to available washers and dryers, attempts assignment of an overweight load, and starts machine cycle timers", "Machine status switches from available to in-use; overweight load is prevented with a capacity warning; cycle duration and usage count update", "Machine status updated to in-use; capacity limit enforced; cycle timers and usage count tracked accurately", "Passed"),
        ("TC006", "Payment Processing", "Customer selects cash or cashless payment methods, and staff updates cash payment status upon counter settlement", "Cash payment is marked pending; cashless payment is marked paid; manual status update marks order as paid", "Cash initialized as pending; cashless payment recorded as paid; manual update changed status to paid", "Passed"),
        ("TC007", "Order Tracking", "Customer tracks order progress across service stages, submits cancellation request, and verifies cancellation through a one-time verification code", "Current order stage and payment status are displayed; invalid verification code is rejected; valid verification code cancels order and releases queue position", "Order stages displayed accurately; invalid verification code rejected; valid verification code cancelled order and cleared queue", "Passed"),
        ("TC008", "Demand Forecasting", "Historical transaction records are processed by the regression model to forecast hourly customer demand", "Peak operating hours are projected; forecast accuracy is measured and recorded against acceptable performance limits", "Peak operating hours projected; forecast accuracy remained within acceptable limits", "Passed"),
        ("TC009", "Dashboard & Auditing", "Administrator accesses dashboard analytics, filters operational records, and performs administrative updates", "Daily statistics, machine utilization, and revenue totals compute correctly; record filters update dynamically; administrative actions are logged", "Dashboard values computed accurately; filter views updated dynamically; administrative actions recorded successfully", "Passed")
    ]
    
    table = doc.add_table(rows=len(table_data) + 2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    col_widths = [Inches(0.75), Inches(1.15), Inches(1.35), Inches(1.50), Inches(1.25), Inches(0.60)]
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=80, right=80)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))
    
    # Subheading row: "Web Application"
    sub_cells = table.rows[1].cells
    sub_cells[0].merge(sub_cells[5])
    sub_cells[0].text = "Web Application"
    p_sub = sub_cells[0].paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sub.paragraph_format.line_spacing = 1.0
    p_sub.paragraph_format.space_after = Pt(2)
    p_sub.paragraph_format.space_before = Pt(2)
    for run in p_sub.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.bold = True
    set_cell_background(sub_cells[0], "F2F2F2")
    set_cell_margins(sub_cells[0], top=50, bottom=50, left=80, right=80)
    
    # Data Rows
    for row_idx, data_tuple in enumerate(table_data, start=2):
        row_cells = table.rows[row_idx].cells
        for col_idx, text in enumerate(data_tuple):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if col_idx in [0, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                if col_idx == 0:
                    run.font.bold = True
            set_cell_margins(row_cells[col_idx], top=50, bottom=50, left=80, right=80)
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w
            
    # Summary paragraph
    add_body_p(
        doc,
        "Table 24 presents the unit and integration testing results for the Web Application. A total of nine (9) "
        "comprehensive test cases were executed across the primary functional components of the system, comprising "
        "Authentication, Order Booking, Order Verification, Queue Management, Machine Management, Payment Processing, "
        "Order Tracking, Demand Forecasting, and Dashboard & Auditing. All nine test cases achieved a status of Passed, "
        "demonstrating that each component operates reliably under both normal conditions and constraint-handling "
        "scenarios. The successful execution confirms that the web application satisfies its design requirements and "
        "provides dependable operational support for the laundry management workflow.",
        space_after=12
    )
    
    # Save outputs safely
    for p in [r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_DEVELOPMENT_AND_TESTING.docx',
              r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_PRESENTATION_OF_RESULTS.docx']:
        try:
            doc.save(p)
            print(f"Successfully saved: {p}")
        except Exception as e:
            print(f"Could not save {p}: {e}")

    try:
        doc.save(r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_RESULTS_AND_DISCUSSION.docx')
        print("Successfully saved: CHAPTER_4_RESULTS_AND_DISCUSSION.docx")
    except Exception:
        print("Note: CHAPTER_4_RESULTS_AND_DISCUSSION.docx is currently open in Word.")

if __name__ == '__main__':
    build_full_grc_chapter4()

