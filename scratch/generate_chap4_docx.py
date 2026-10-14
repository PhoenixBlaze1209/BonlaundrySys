import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_chapter4_doc(output_path):
    doc = docx.Document()

    # Page Setup: Standard Letter with 1-inch margins
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    # Chapter Header
    p_chap = doc.add_paragraph()
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_chap.paragraph_format.space_after = Pt(2)
    p_chap.paragraph_format.line_spacing = 1.0
    r_chap = p_chap.add_run("CHAPTER IV")
    r_chap.bold = True
    r_chap.font.size = Pt(14)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(18)
    p_title.paragraph_format.line_spacing = 1.0
    r_title = p_title.add_run("RESULTS AND DISCUSSION")
    r_title.bold = True
    r_title.font.size = Pt(14)

    # Introductory Paragraph
    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.space_after = Pt(14)
    p_intro.add_run(
        "This chapter presents the results and discussions of the Web-Based Laundry Management "
        "System with Machine Learning for Demand Forecasting and Fleet Optimization developed for "
        "Bon Laundry Station. The chapter is structured in accordance with the research methodology "
        "and follows the prescribed outline: Development and Testing, Evaluation Results, and Data "
        "Interpretation. It covers the outcomes of unit and integration testing, the results of the "
        "ISO/IEC 25010-based user acceptance testing conducted with IT Experts and End Users, and the "
        "interpretation of the system's overall acceptability. The findings demonstrate how the "
        "system's web-based admin dashboard and customer portal address the specific objectives of "
        "automating booking, managing laundry queues and machine allocations, tracking service progress, "
        "and providing predictive demand analytics."
    )

    # Section 4.1 Header
    p_s41 = doc.add_paragraph()
    p_s41.paragraph_format.space_before = Pt(8)
    p_s41.paragraph_format.space_after = Pt(4)
    p_s41.paragraph_format.line_spacing = 1.15
    r_s41 = p_s41.add_run("4.1 Development and Testing")
    r_s41.bold = True
    r_s41.font.size = Pt(12)

    # Subsection: Unit and Integration Testing
    p_uit = doc.add_paragraph()
    p_uit.paragraph_format.space_before = Pt(4)
    p_uit.paragraph_format.space_after = Pt(8)
    p_uit.paragraph_format.line_spacing = 1.15
    r_uit = p_uit.add_run("Unit and Integration Testing")
    r_uit.bold = True
    r_uit.font.size = Pt(12)

    # Table Title
    p_tbl_title = doc.add_paragraph()
    p_tbl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_title.paragraph_format.space_before = Pt(8)
    p_tbl_title.paragraph_format.space_after = Pt(6)
    r_tbl_title = p_tbl_title.add_run("Table 24 Unit and Integration Testing for Web Application")
    r_tbl_title.italic = True
    r_tbl_title.font.size = Pt(11)

    # Table Setup
    # 6 columns: Test Case ID, Component/Module, Test Condition, Expected Behavior, Actual Result, Status
    headers = [
        "Test\nCase\nID",
        "Component/Module",
        "Test\nCondition",
        "Expected\nBehavior",
        "Actual Result",
        "Status"
    ]

    table_data = [
        (
            "TC001",
            "Authentication",
            "User logs in with valid credentials, attempts login with incorrect password, and logs in using single sign-on authentication",
            "Valid credentials grant access to the appropriate role interface; invalid login is refused with no session created; single sign-on authenticates and redirects user",
            "Access granted for valid credentials; access refused for invalid credentials; single sign-on authenticated and redirected successfully",
            "Passed"
        ),
        (
            "TC002",
            "Order Booking",
            "Customer places a laundry booking with selected service package, specifies condominium door-to-door delivery with missing and complete address details",
            "Incomplete address details are blocked with validation prompts; valid booking is saved with calculated service charges and delivery fees",
            "Missing address fields blocked; valid booking accepted and charges calculated correctly",
            "Passed"
        ),
        (
            "TC003",
            "Order Verification",
            "Staff searches for an order using an order code, attempts lookup of an invalid code, and confirms garment intake",
            "Valid order details are retrieved and displayed; invalid code is rejected with an error notice; order confirmation updates the order status and moves it to the queue",
            "Order details retrieved correctly; invalid code rejected; order verified and moved to active queue",
            "Passed"
        ),
        (
            "TC004",
            "Queue Management",
            "Verified orders enter the active queue, waiting times are calculated, and orders transition to processing upon machine allocation",
            "Queue numbers are issued sequentially; estimated waiting times compute based on machine load; queue status updates as orders progress",
            "Queue numbers issued in order; waiting times calculated accurately; status transitions updated without delay",
            "Passed"
        ),
        (
            "TC005",
            "Machine Management",
            "Staff assigns laundry loads to available washers and dryers, attempts assignment of an overweight load, and starts machine cycle timers",
            "Machine status switches from available to in-use; overweight load is prevented with a capacity warning; cycle duration and usage count update",
            "Machine status updated to in-use; capacity limit enforced; cycle timers and usage count tracked accurately",
            "Passed"
        ),
        (
            "TC006",
            "Payment Processing",
            "Customer selects cash or cashless payment methods, and staff updates cash payment status upon counter settlement",
            "Cash payment is marked pending; cashless payment is marked paid; manual status update marks order as paid",
            "Cash initialized as pending; cashless payment recorded as paid; manual update changed status to paid",
            "Passed"
        ),
        (
            "TC007",
            "Order Tracking",
            "Customer tracks order progress across service stages, submits cancellation request, and verifies cancellation through a one-time verification code",
            "Current order stage and payment status are displayed; invalid verification code is rejected; valid verification code cancels order and releases queue position",
            "Order stages displayed accurately; invalid verification code rejected; valid verification code cancelled order and cleared queue",
            "Passed"
        ),
        (
            "TC008",
            "Demand Forecasting",
            "Historical transaction records are processed by the regression model to forecast hourly customer demand",
            "Peak operating hours are projected; forecast accuracy is measured and recorded against acceptable performance limits",
            "Peak operating hours projected; forecast accuracy remained within acceptable limits",
            "Passed"
        ),
        (
            "TC009",
            "Dashboard & Auditing",
            "Administrator accesses dashboard analytics, filters operational records, and performs administrative updates",
            "Daily statistics, machine utilization, and revenue totals compute correctly; record filters update dynamically; administrative actions are logged",
            "Dashboard values computed accurately; filter views updated dynamically; administrative actions recorded successfully",
            "Passed"
        )
    ]

    # Create table: 1 header row + 1 subheader row + 9 data rows = 11 rows
    table = doc.add_table(rows=len(table_data) + 2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    col_widths = [Inches(0.75), Inches(1.15), Inches(1.35), Inches(1.50), Inches(1.25), Inches(0.60)]

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Repeat header row on every page
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    # Subheading row: "Web Application"
    sub_cells = table.rows[1].cells
    sub_cells[0].merge(sub_cells[5])
    sub_cells[0].text = "Web Application"
    p_sub = sub_cells[0].paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sub.paragraph_format.line_spacing = 1.0
    p_sub.paragraph_format.space_after = Pt(3)
    p_sub.paragraph_format.space_before = Pt(3)
    for run in p_sub.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.bold = True
    set_cell_background(sub_cells[0], "F2F2F2")
    set_cell_margins(sub_cells[0], top=60, bottom=60, left=100, right=100)

    # Data Rows
    for row_idx, data_tuple in enumerate(table_data, start=2):
        row_cells = table.rows[row_idx].cells
        for col_idx, text in enumerate(data_tuple):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if col_idx in [0, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                if col_idx == 0:
                    run.font.bold = True
            set_cell_margins(row_cells[col_idx], top=60, bottom=60, left=90, right=90)
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # Set column widths for all rows
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # Concluding Analysis Paragraph
    p_summary = doc.add_paragraph()
    p_summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_summary.paragraph_format.first_line_indent = Inches(0.5)
    p_summary.paragraph_format.line_spacing = 1.5
    p_summary.paragraph_format.space_before = Pt(14)
    p_summary.paragraph_format.space_after = Pt(12)
    p_summary.add_run(
        "Table 24 presents the unit and integration testing results for the Web Application. "
        "A total of nine (9) comprehensive test cases were executed across the primary functional "
        "components of the system, comprising Authentication, Order Booking, Order Verification, "
        "Queue Management, Machine Management, Payment Processing, Order Tracking, Demand Forecasting, "
        "and Dashboard & Auditing. All nine test cases achieved a status of Passed, demonstrating "
        "that each component operates reliably under both normal conditions and constraint-handling "
        "scenarios. The successful execution confirms that the web application satisfies its design "
        "requirements and provides dependable operational support for the laundry management workflow."
    )

    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == '__main__':
    build_chapter4_doc(r'c:\Users\genes\bon-laundry\docs\CHAPTER_4_DEVELOPMENT_AND_TESTING.docx')
