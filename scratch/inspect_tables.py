import docx

doc = docx.Document('docs/BON_LAUNDRY_CUSTOMER_PORTAL_FUNCTIONAL_SPECIFICATION.docx')
with open('scratch/tables_summary.txt', 'w', encoding='utf-8') as f:
    f.write(f"Sections: {len(doc.sections)}\n")
    for si, sec in enumerate(doc.sections):
        f.write(f"Section {si}: w={sec.page_width.inches}, h={sec.page_height.inches}, top={sec.top_margin.inches}, bot={sec.bottom_margin.inches}, left={sec.left_margin.inches}, right={sec.right_margin.inches}\n")
    
    f.write(f"\nTables: {len(doc.tables)}\n")
    for ti, t in enumerate(doc.tables):
        f.write(f"\n--- Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) ---\n")
        for ri, r in enumerate(t.rows):
            row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
            f.write(f"  Row {ri}: {' | '.join(row_txt)}\n")
print("Done writing scratch/tables_summary.txt")
