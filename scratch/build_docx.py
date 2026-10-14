# scratch/build_docx.py
"""
Converts docs/USER_MANUAL.md into a high-polish, professional Microsoft Word document (docs/USER_MANUAL.docx).
Includes custom styling, typography, embedded screenshots, shaded tables, callout alert boxes,
code blocks, and formatted headers/footers.
"""
import os
import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
MD_PATH = BASE_DIR / 'docs' / 'USER_MANUAL.md'
DOCX_PATH = BASE_DIR / 'docs' / 'USER_MANUAL.docx'
SCREENSHOTS_DIR = BASE_DIR / 'docs' / 'screenshots'

# Palette Constants
COLOR_PRIMARY = RGBColor(19, 37, 49)      # Service Ink #132531
COLOR_ACCENT = RGBColor(2, 132, 199)      # Brand Blue #0284C7
COLOR_SECONDARY = RGBColor(51, 65, 85)    # Slate #334155
COLOR_BODY = RGBColor(30, 41, 59)         # Ink Body #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Gray #64748B
COLOR_LIGHT_BG = "F8FAFC"
COLOR_HEADER_BG = "0284C7"
COLOR_CALLOUT_BG = "F0F9FF"
COLOR_WARNING_BG = "FEF3C7"
COLOR_BORDER = "CBD5E1"


def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell margins (padding) in dxa (1/20 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Applies subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout(doc, text, alert_type="NOTE"):
    """Adds a callout box with a colored left border and light background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)

    is_warning = alert_type in ("WARNING", "CAUTION")
    border_color = "F59E0B" if is_warning else "0284C7"
    bg_color = COLOR_WARNING_BG if is_warning else COLOR_CALLOUT_BG

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    # Prefix
    run_icon = p.add_run(f"[{alert_type}] ")
    run_icon.bold = True
    run_icon.font.name = "Arial"
    run_icon.font.size = Pt(10)
    run_icon.font.color.rgb = RGBColor(146, 64, 14) if is_warning else COLOR_ACCENT

    # Text content
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = COLOR_BODY

    # Empty paragraph for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def format_runs(p, text):
    """Parses basic inline markdown: **bold**, *italic*, `code`, math $$, and plain text."""
    pattern = r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$\$[^\$]+?\$\$|\$[^\$]+?\$|\[[^\]]+?\]\([^\)]+?\))'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.color.rgb = COLOR_BODY
        elif token.startswith('*') and token.endswith('*') and len(token) > 1:
            run = p.add_run(token[1:-1])
            run.italic = True
            run.font.color.rgb = COLOR_BODY
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 59, 64)
        elif (token.startswith('$$') and token.endswith('$$')) or (token.startswith('$') and token.endswith('$')):
            # Math formula representation
            clean_math = token.strip('$').replace(r'\xrightarrow{\text{inherits}}', ' ──▶ (inherits) ──▶ ').replace(r'\text{', '').replace('}', '')
            run = p.add_run(clean_math)
            run.bold = True
            run.font.name = "Consolas"
            run.font.color.rgb = COLOR_ACCENT
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            link_text = token[1:token.index('](')]
            run = p.add_run(link_text)
            run.font.color.rgb = COLOR_ACCENT
            run.underline = True
        else:
            run = p.add_run(token)
            run.font.color.rgb = COLOR_BODY
        run.font.name = "Arial"


def compile_docx():
    print(f"Reading markdown from: {MD_PATH}")
    md_content = MD_PATH.read_text(encoding='utf-8')
    lines = md_content.splitlines()

    doc = Document()

    # Configure Margins: 1 inch (72pt)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Configure Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Bon Laundry Station — System User Manual")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun_left = fp.add_run("Urban Deca Homes Ortigas Branch  •  Confidential Operations Manual")
        frun_left.font.name = "Arial"
        frun_left.font.size = Pt(8.5)
        frun_left.font.color.rgb = COLOR_MUTED

    # Style defaults
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)

    in_code_block = False
    code_block_lang = ""
    code_lines = []

    in_table = False
    table_lines = []

    in_callout = False
    callout_type = "NOTE"
    callout_lines = []

    i = 0
    total_lines = len(lines)

    while i < total_lines:
        line = lines[i]

        # 1. Code Block boundary
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line[3:].strip()
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                # Render code block / diagram
                if code_lines:
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = tbl.cell(0, 0)
                    cell.width = Inches(6.5)
                    set_cell_background(cell, "F1F5F9")
                    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                    
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcBorders = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                        f'<w:left w:val="single" w:sz="16" w:space="0" w:color="0284C7"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                        f'</w:tcBorders>'
                    )
                    tcPr.append(tcBorders)

                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.line_spacing = 1.15
                    
                    code_text = "\n".join(code_lines)
                    crun = cp.add_run(code_text)
                    crun.font.name = "Consolas"
                    crun.font.size = Pt(8.5)
                    crun.font.color.rgb = COLOR_BODY
                    
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_before = Pt(0)
                    p_spacer.paragraph_format.space_after = Pt(4)

                code_lines = []
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 2. Callout / Alert block
        if line.startswith('> [!'):
            m = re.match(r'> \[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]', line)
            if m:
                alert_type = m.group(1)
                callout_text_lines = []
                j = i + 1
                while j < total_lines and lines[j].startswith('>'):
                    cleaned = lines[j][1:].strip()
                    if cleaned:
                        callout_text_lines.append(cleaned)
                    j += 1
                add_callout(doc, " ".join(callout_text_lines), alert_type)
                i = j
                continue

        # 3. Markdown Tables
        if line.startswith('|') and '|' in line[1:]:
            table_lines.append(line)
            j = i + 1
            while j < total_lines and lines[j].strip().startswith('|'):
                table_lines.append(lines[j].strip())
                j += 1

            # Parse and render table
            raw_rows = []
            for tl in table_lines:
                # skip divider rows like |---|---|
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                if cells and all(set(c).issubset({'-', ':', ' '}) for c in cells):
                    continue
                raw_rows.append(cells)

            if raw_rows:
                num_cols = max(len(r) for r in raw_rows)
                tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="4")

                for row_idx, rdata in enumerate(raw_rows):
                    is_header = (row_idx == 0)
                    row_el = tbl.rows[row_idx]
                    
                    # Prevent row split across pages
                    trPr = row_el._tr.get_or_add_trPr()
                    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                    if is_header:
                        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

                    for col_idx in range(num_cols):
                        cell = row_el.cells[col_idx]
                        text_val = rdata[col_idx] if col_idx < len(rdata) else ""
                        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

                        if is_header:
                            set_cell_background(cell, COLOR_HEADER_BG)
                        elif row_idx % 2 == 1:
                            set_cell_background(cell, COLOR_LIGHT_BG)
                        else:
                            set_cell_background(cell, "FFFFFF")

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.line_spacing = 1.1

                        if is_header:
                            run = p.add_run(text_val)
                            run.bold = True
                            run.font.name = "Arial"
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            format_runs(p, text_val)
                            for r in p.runs:
                                r.font.size = Pt(9)

                p_after_tbl = doc.add_paragraph()
                p_after_tbl.paragraph_format.space_before = Pt(2)
                p_after_tbl.paragraph_format.space_after = Pt(6)

            table_lines = []
            i = j
            continue

        # 4. Images
        img_match = re.match(r'!\[(.*?)\]\((file:///[^)]+)\)', line)
        if img_match:
            caption = img_match.group(1)
            raw_url = img_match.group(2)
            # resolve path
            file_path_str = raw_url.replace('file:///', '').replace('/', os.sep)
            img_file = Path(file_path_str)
            if not img_file.exists():
                # check inside SCREENSHOTS_DIR
                img_file = SCREENSHOTS_DIR / img_file.name

            if img_file.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                doc.add_picture(str(img_file), width=Inches(5.8))

                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(12)
                c_run = p_cap.add_run(caption)
                c_run.font.name = "Arial"
                c_run.font.size = Pt(9)
                c_run.italic = True
                c_run.font.color.rgb = COLOR_MUTED

            i += 1
            continue

        # 5. Headings
        if line.startswith('# '):
            h_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(22)
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(15)
            run.font.color.rgb = COLOR_ACCENT
            i += 1
            continue

        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        if line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # 6. Horizontal Rules
        if line.strip() in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # 7. Bullet Lists & Numbered Lists
        bullet_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)', line)
        if bullet_match:
            indent_space, marker, item_text = bullet_match.groups()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            # Indent level
            indent_level = len(indent_space) // 2
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))

            bullet_run = p.add_run("• " if marker in ('-', '*') else f"{marker} ")
            bullet_run.bold = True
            bullet_run.font.name = "Arial"
            bullet_run.font.size = Pt(10)
            bullet_run.font.color.rgb = COLOR_ACCENT

            format_runs(p, item_text)
            i += 1
            continue

        # 8. Regular Paragraphs
        stripped = line.strip()
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            format_runs(p, stripped)

        i += 1

    alias_path = BASE_DIR / 'docs' / 'BON_LAUNDRY_USER_MANUAL.docx'
    print(f"Saving compiled DOCX to: {alias_path}")
    doc.save(str(alias_path))
    print(f"Saved alias: {alias_path.name} ({alias_path.stat().st_size / 1024:.1f} KB)")

    try:
        print(f"Saving compiled DOCX to: {DOCX_PATH}")
        doc.save(str(DOCX_PATH))
        print(f"Compilation Complete! Size: {DOCX_PATH.stat().st_size / 1024:.1f} KB")
    except PermissionError:
        fallback_path = BASE_DIR / 'docs' / 'USER_MANUAL_UPDATED.docx'
        doc.save(str(fallback_path))
        print(f"Note: USER_MANUAL.docx was locked by Word. Saved to: {fallback_path.name} ({fallback_path.stat().st_size / 1024:.1f} KB)")


if __name__ == '__main__':
    compile_docx()
